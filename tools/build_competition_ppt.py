"""生成“田诊卫士 · 智能体版”比赛答辩 PPT（14 页，16:9）。

用法：python tools/build_competition_ppt.py
输出：presentation/田诊卫士-智能体版-答辩.pptx
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentation" / "田诊卫士-智能体版-答辩.pptx"
BUILD = ROOT / ".build" / "competition-ppt"
SHOTS = ROOT / "reports" / "screenshots"
FIGURES = ROOT / "app" / "figures"

# 调色板：深绿为主，浅绿为底，琥珀色只用于关键提醒。
BG = "F3F7F3"
WHITE = "FFFFFF"
DARK = "0E3B2E"
PRIMARY = "124F3D"
GREEN = "146C50"
LIGHT = "E6F1E8"
LIGHT_2 = "F7FAF7"
AMBER = "D9A441"
AMBER_BG = "FFF6DF"
AMBER_INK = "875A12"
RED = "9D3529"
RED_BG = "F8E2DF"
INK = "18372E"
MUTED = "61766D"
LINE = "DCE7DF"

FONT = "微软雅黑"
W, H = 13.333, 7.5
M = 0.55


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _set_font(run, size: float, bold: bool, color: str, font: str = FONT) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", font)


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    content: str,
    size: float = 14,
    bold: bool = False,
    color: str = INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing: float = 1.08,
    space_after: float = 3,
    font: str = FONT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = str(content).split("\n")
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = line
        _set_font(run, size, bold, color, font)
    return box


def bullets(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items: list[str],
    size: float = 13,
    color: str = INK,
    bullet_color: str = GREEN,
    line_spacing: float = 1.12,
    gap: float = 7,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(gap)
        ppr = paragraph._p.get_or_add_pPr()
        ppr.set("marL", "171450")
        ppr.set("indent", "-171450")
        run = paragraph.add_run()
        run.text = "▪  "
        _set_font(run, size, True, bullet_color)
        run2 = paragraph.add_run()
        run2.text = item
        _set_font(run2, size, False, color)
    return box


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str | None = WHITE,
    line: str | None = LINE,
    radius: float | None = 0.08,
    line_width: float = 0.75,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = radius
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, width: float = 1.4):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    ln = connector._element.spPr.ln
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return connector


def bg(slide, color: str = BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def title(slide, kicker: str, heading: str, color: str = PRIMARY, kicker_color: str = GREEN):
    text(slide, M, 0.34, 12.2, 0.3, kicker, 10.5, True, kicker_color)
    text(slide, M, 0.63, 12.2, 0.7, heading, 27, True, color)


def footer(slide, page: int, dark: bool = False):
    color = "A8D4BF" if dark else MUTED
    text(slide, M, 7.12, 7.0, 0.25, "田诊卫士 · 智能体版 · 智慧农业", 9, False, color)
    text(slide, 11.4, 7.12, 1.4, 0.25, f"{page:02d} / 14", 9, False, color, PP_ALIGN.RIGHT)


def rounded_image(source: Path, target: Path, radius: int = 20) -> Path:
    BUILD.mkdir(parents=True, exist_ok=True)
    image = Image.open(source).convert("RGBA")
    mask = Image.new("L", image.size, 0)
    draw = ImageDraw.Draw(mask)
    draw.rounded_rectangle([0, 0, image.width, image.height], radius=radius, fill=255)
    image.putalpha(mask)
    image.save(target)
    return target


def picture_card(slide, image_path: Path, x: float, y: float, w: float, pad: float = 0.09, caption: str | None = None):
    image = Image.open(image_path)
    h = w * image.height / image.width
    rect(slide, x - pad, y - pad, w + 2 * pad, h + 2 * pad, WHITE, LINE, 0.04)
    slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if caption:
        text(slide, x, y + h + 0.12, w, 0.28, caption, 9.5, False, MUTED)
    return h


def add_table(
    slide,
    x: float,
    y: float,
    w: float,
    rows: list[list[str]],
    col_widths: list[float],
    row_height: float = 0.45,
    font_size: float = 10.5,
    header_size: float = 10.5,
    header_fill: str = LIGHT,
):
    nrows, ncols = len(rows), len(rows[0])
    height = row_height * nrows
    shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(height))
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    for col, width in enumerate(col_widths):
        table.columns[col].width = Inches(width)
    for row_idx, row in enumerate(rows):
        table.rows[row_idx].height = Inches(row_height)
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if row_idx == 0 else (WHITE if row_idx % 2 else LIGHT_2))
            tf = cell.text_frame
            tf.word_wrap = True
            paragraph = tf.paragraphs[0]
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = value
            _set_font(
                run,
                header_size if row_idx == 0 else font_size,
                row_idx == 0,
                PRIMARY if row_idx == 0 else INK,
            )
    return table


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # ---------------------------------------------------------------- 01 封面
    slide = prs.slides.add_slide(blank)
    bg(slide, DARK)
    text(slide, 0.8, 0.7, 6.0, 0.3, "智慧农业 · 智能体创新应用比赛", 12, True, "A8D4BF")
    text(slide, 0.8, 1.12, 6.0, 0.9, "田诊卫士", 46, True, WHITE)
    text(slide, 0.8, 2.05, 6.0, 0.8, "面向小农户的番茄病害诊断\n与安全用药智能体", 20, False, "D7E8DD", line_spacing=1.18)
    text(slide, 0.8, 3.15, 6.1, 0.4, "大模型规划  ·  工具调用  ·  规则否决", 14, True, AMBER)
    text(
        slide,
        0.8,
        3.7,
        5.7,
        1.2,
        "从一句口语描述开始，调用视觉诊断、病害检索、候选记录、规则校验和施药窗口工具，给出可回查的依据与缺口。",
        12.5,
        False,
        "B8D4C5",
        line_spacing=1.35,
    )
    text(slide, 0.8, 6.25, 6.4, 0.3, "“码上智农，电亮未来”智能体创新应用比赛", 10.5, False, "A8D4BF")
    text(slide, 0.8, 6.58, 6.4, 0.3, "团队：【队名】  ·  2026 年 10 月", 10.5, False, "A8D4BF")

    cover_img = rounded_image(SHOTS / "agent-02-回答结论.png", BUILD / "cover.png", 22)
    picture_card(slide, cover_img, 7.35, 1.35, 5.45)
    text(slide, 7.35, 4.72, 5.45, 0.3, "真实运行界面：对话 + 工具轨迹 + 逐条校验", 9.5, False, "A8D4BF")

    # ---------------------------------------------------------------- 02 问题
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "01 · 问题", "识别出病名，不等于能安全用药")
    text(slide, M, 1.6, 5.0, 0.5, "真正的风险，发生在识别之后。", 22, True, GREEN)
    text(
        slide,
        M,
        2.2,
        5.0,
        0.8,
        "小农户拿到一个病名之后，还要完成一整套决策：查登记、核剂量、算间隔期、看次数、等窗口。",
        12.5,
        False,
        MUTED,
        line_spacing=1.35,
    )
    bullets(
        slide,
        M,
        3.25,
        5.0,
        2.2,
        [
            "登记作物和防治对象是否匹配？",
            "剂量、单位、安全间隔期是否满足？",
            "本季次数是否超限，什么时候能打？",
        ],
        13,
    )
    rect(slide, M, 5.55, 5.0, 1.1, AMBER_BG, None, 0.08)
    text(slide, M + 0.18, 5.7, 4.65, 0.8, "5 天后采收 + 15 天安全间隔期\n→ 代森锰锌被直接拦截", 13, True, AMBER_INK, line_spacing=1.3)

    steps = [
        ("1", "农户描述 / 拍照", "口语症状、叶片照片、采收计划"),
        ("2", "智能体理解与追问", "信息不足先澄清，不猜病名"),
        ("3", "调用工具", "诊断、检索、候选、校验、窗口"),
        ("4", "规则否决与回答", "通过 / 待核验 / 拦截，附来源和缺口"),
    ]
    arrow_x = 6.33
    text(slide, 6.05, 1.55, 6.5, 0.3, "一次问诊的完整链路", 10.5, True, GREEN)
    for index, (number, heading, description) in enumerate(steps):
        y = 1.95 + index * 1.2
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.05), Inches(y + 0.18), Inches(0.52), Inches(0.52))
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(PRIMARY if index < 3 else AMBER)
        circle.line.fill.background()
        circle.shadow.inherit = False
        text(slide, 6.05, y + 0.245, 0.52, 0.3, number, 11, True, WHITE, PP_ALIGN.CENTER)
        rect(slide, 6.75, y, 5.95, 0.92, WHITE, LINE, 0.08)
        text(slide, 6.92, y + 0.1, 5.6, 0.3, heading, 12.5, True, PRIMARY)
        text(slide, 6.92, y + 0.43, 5.6, 0.3, description, 10, False, MUTED)
        if index < len(steps) - 1:
            arrow(slide, 6.31, y + 0.78, 6.31, y + 1.12, "CFE0D5", 1.2)
    footer(slide, 2)

    # ---------------------------------------------------------------- 03 连续决策
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "02 · 场景", "小农户的真实问题，是一个连续决策")
    cards = [
        ("1", "看到症状", "叶背白毛、茎上黑斑，描述模糊"),
        ("2", "判断病害", "像晚疫病，但不能确诊"),
        ("3", "查药核对", "登记？剂量？间隔期？次数？"),
        ("4", "决定何时打", "天气、采收期、可用窗口"),
    ]
    for index, (number, heading, description) in enumerate(cards):
        x = M + index * 3.2
        rect(slide, x, 1.65, 2.95, 1.65, WHITE, LINE, 0.08)
        text(slide, x + 0.18, 1.78, 0.5, 0.4, number, 18, True, AMBER)
        text(slide, x + 0.18, 2.2, 2.55, 0.35, heading, 13, True, PRIMARY)
        text(slide, x + 0.18, 2.62, 2.55, 0.5, description, 10, False, MUTED, line_spacing=1.25)
        if index < 3:
            arrow(slide, x + 2.99, 2.45, x + 3.15, 2.45, "A9C9B7", 1.3)

    rect(slide, M, 3.95, 6.0, 2.35, LIGHT, None, 0.08)
    text(slide, M + 0.22, 4.1, 5.5, 0.35, "我们做", 15, True, GREEN)
    bullets(
        slide,
        M + 0.22,
        4.52,
        5.5,
        1.6,
        ["理解口语和照片，先判断信息够不够", "检索知识库并逐条做确定性校验", "展示来源、核验状态和模拟窗口", "把缺失信息问出来，而不是补猜"],
        11.5,
        gap=5,
    )
    rect(slide, 6.8, 3.95, 6.0, 2.35, AMBER_BG, None, 0.08)
    text(slide, 7.02, 4.1, 5.5, 0.35, "我们不做", 15, True, AMBER_INK)
    bullets(
        slide,
        7.02,
        4.52,
        5.5,
        1.6,
        ["不替代农技员做最终诊断", "不签发正式田间处方", "不把模拟天气说成当地预报", "不把待核验数据说成已通过"],
        11.5,
        color=AMBER_INK,
        bullet_color=AMBER_INK,
        gap=5,
    )
    text(slide, M, 6.55, 12.2, 0.3, "目标：把依据和缺口摆到桌面上，而不是给一个看似确定的答案。", 11, False, MUTED)
    footer(slide, 3)

    # ---------------------------------------------------------------- 04 架构
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "03 · 架构", "智能体架构：大模型规划，规则引擎否决")

    rect(slide, M, 2.75, 1.75, 1.25, LIGHT, None, 0.1)
    text(slide, M + 0.1, 2.95, 1.55, 0.8, "用户\n口语 / 照片", 12, True, PRIMARY, PP_ALIGN.CENTER, line_spacing=1.3)
    arrow(slide, 2.32, 3.38, 2.92, 3.38, GREEN, 1.6)

    rect(slide, 2.95, 2.3, 3.3, 2.2, PRIMARY, None, 0.08)
    text(slide, 3.12, 2.48, 2.95, 0.35, "大模型规划器", 17, True, WHITE)
    text(slide, 3.12, 2.92, 2.95, 1.3, "理解 · 追问 · 选择工具\n\n不直接生成农药剂量、\n间隔期或登记证号", 10.5, False, "D7E8DD", line_spacing=1.35)

    tool_names = [
        ("get_leaf_diagnosis", "视觉诊断"),
        ("search_disease", "症状检索"),
        ("get_candidate_records", "候选记录"),
        ("validate_prescription", "规则校验"),
        ("get_spray_window", "施药窗口"),
        ("ask_user", "追问澄清"),
        ("get_kb_status", "知识库状态"),
    ]
    text(slide, 7.0, 1.55, 2.55, 0.3, "工具（function calling）", 10.5, True, GREEN)
    for index, (name, label) in enumerate(tool_names):
        y = 1.78 + index * 0.56
        rect(slide, 7.0, y, 2.55, 0.46, WHITE, LINE, 0.12)
        text(slide, 7.12, y + 0.05, 1.55, 0.2, name, 8.6, True, PRIMARY)
        text(slide, 8.55, y + 0.05, 0.9, 0.2, label, 8.3, False, MUTED, PP_ALIGN.RIGHT)
    arrow(slide, 6.28, 3.38, 6.92, 3.38, GREEN, 1.6)
    text(slide, 6.28, 3.1, 0.7, 0.22, "tool calls", 8.5, False, MUTED, PP_ALIGN.CENTER)

    rect(slide, 2.95, 5.1, 3.3, 1.1, AMBER_BG, None, 0.08)
    text(slide, 3.12, 5.22, 2.95, 0.3, "确定性规则引擎", 14, True, AMBER_INK)
    text(slide, 3.12, 5.57, 2.95, 0.4, "8 条规则 · 逐条裁决", 10.5, False, AMBER_INK)
    arrow(slide, 4.6, 4.55, 4.6, 5.05, MUTED, 1.5)

    rect(slide, 7.0, 5.8, 5.75, 0.95, PRIMARY, None, 0.08)
    text(slide, 7.18, 5.9, 5.4, 0.75, "通过 / 待核验 / 拦截\n+ 来源 + 模拟窗口 + 安全边界", 11.5, True, WHITE, line_spacing=1.3)
    arrow(slide, 6.28, 5.7, 6.92, 5.7, MUTED, 1.5)

    text(slide, M, 6.55, 6.0, 0.3, "模型负责理解与规划；代码负责事实与否决。", 11, False, MUTED)
    footer(slide, 4)

    # ---------------------------------------------------------------- 05 完整对话
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "04 · 演示", "一次完整对话：从症状到逐条裁决")
    convo = rounded_image(SHOTS / "agent-02-回答结论.png", BUILD / "convo.png", 18)
    picture_card(slide, convo, M, 1.5, 7.05)

    steps = [
        ("1", "检索", "晚疫病 9.04；灰霉 3.46；早疫 2.50"),
        ("2", "候选", "5 条知识库记录，全部待核验"),
        ("3", "校验", "代森锰锌拦截；其余 4 条待补充"),
        ("4", "窗口", "10/21 06:00–08:00；最晚 10/22"),
        ("5", "回答", "结论、依据、待补信息、边界"),
    ]
    for index, (number, heading, description) in enumerate(steps):
        y = 1.55 + index * 0.93
        text(slide, 8.0, y, 0.35, 0.35, number, 14, True, AMBER)
        text(slide, 8.38, y - 0.02, 4.35, 0.3, heading, 12.5, True, PRIMARY)
        text(slide, 8.38, y + 0.3, 4.35, 0.35, description, 10, False, MUTED, line_spacing=1.25)
        if index < 4:
            text(slide, 8.18, y + 0.62, 0.2, 0.2, "│", 9, False, LINE)
    text(slide, M, 6.62, 12.2, 0.3, "截图来自真实 DeepSeek 运行：工具调用 5 次，回答保留来源和未核验原因。", 9.5, False, MUTED)
    footer(slide, 5)

    # ---------------------------------------------------------------- 06 工具轨迹
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "05 · 工具链", "工具调用轨迹：每一次回答都能回查")
    trace = rounded_image(SHOTS / "agent-03-工具调用.png", BUILD / "trace.png", 18)
    picture_card(slide, trace, M, 1.5, 7.05)

    rows = [
        ["工具", "作用"],
        ["get_leaf_diagnosis", "读取叶片视觉诊断"],
        ["search_disease", "口语症状 → 候选病害"],
        ["get_candidate_records", "读取药剂记录"],
        ["validate_prescription", "逐条规则校验"],
        ["get_spray_window", "按药剂算窗口"],
        ["ask_user", "信息不足时追问"],
        ["get_kb_status", "知识库覆盖范围"],
    ]
    add_table(slide, 8.0, 1.55, 4.75, rows, [2.05, 2.7], row_height=0.47, font_size=9.3, header_size=9.8)
    text(slide, 8.0, 5.55, 4.75, 0.8, "工具结果是事实来源；模型输出是解释层。\n每次回答都能看到：调用了什么、用了哪条记录、哪条规则拦截。", 10, False, MUTED, line_spacing=1.3)
    footer(slide, 6)

    # ---------------------------------------------------------------- 07 视觉诊断
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "06 · 视觉诊断", "先拒识，再分类：不让无效图片进入用药流程")
    chart = FIGURES / "fig1_confusion_matrix.png"
    picture_card(slide, chart, M, 1.55, 5.45, caption="冻结测试集 1088 张 · 混淆矩阵")

    stats = [
        ("96.88%", "总体准确率"),
        ("6.53%", "测试集拒识率"),
        ("98.82%", "接受样本准确率"),
        ("0.941", "验证集确定阈值"),
    ]
    for index, (number, label) in enumerate(stats):
        x = 6.55 + (index % 2) * 3.15
        y = 1.75 + (index // 2) * 1.75
        text(slide, x, y, 2.9, 0.6, number, 30, True, GREEN)
        text(slide, x, y + 0.62, 2.9, 0.3, label, 11, False, MUTED)
    rect(slide, 6.55, 5.25, 6.1, 1.05, LIGHT, None, 0.08)
    text(slide, 6.73, 5.38, 5.75, 0.8, "纯色、低纹理、严重模糊或过小图片在推理前拒绝；只覆盖番茄四类，非叶片仍需独立评测。", 10, False, GREEN, line_spacing=1.3)
    text(slide, M, 6.62, 12.2, 0.3, "来源：PlantVillage 公开数据集；阈值由验证集确定，测试集不参与选择。", 9, False, MUTED)
    footer(slide, 7)

    # ---------------------------------------------------------------- 08 检索
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "07 · 理解", "口语症状检索：把‘像什么’和‘是什么’分开")
    flow = ["口语描述", "否定片段剔除", "别名 + 症状加权 + BM25", "候选病害 + 检索分"]
    for index, label in enumerate(flow):
        y = 1.65 + index * 1.12
        rect(slide, M, y, 5.5, 0.72, WHITE if index < 3 else LIGHT, LINE, 0.1)
        text(slide, M + 0.2, y + 0.18, 5.1, 0.35, label, 12, index == 3, PRIMARY if index == 3 else INK)
        if index < 3:
            arrow(slide, 3.25, y + 0.75, 3.25, y + 1.08, "A9C9B7", 1.3)
    text(slide, M, 6.2, 5.5, 0.3, "检索分只表示文字相关性，不是确诊概率。", 10.5, True, AMBER_INK)

    rect(slide, 6.7, 1.6, 6.05, 2.0, LIGHT, None, 0.08)
    text(slide, 6.92, 1.76, 5.6, 0.3, "有明确症状", 12, True, GREEN)
    text(slide, 6.92, 2.15, 5.6, 0.3, "“叶子背面长白毛，茎上有黑斑”", 13, True, PRIMARY)
    text(slide, 6.92, 2.6, 5.6, 0.3, "→ 晚疫病 9.04 · 灰霉 3.46 · 早疫 2.50", 11, False, MUTED)
    rect(slide, 6.7, 3.85, 6.05, 2.0, AMBER_BG, None, 0.08)
    text(slide, 6.92, 4.02, 5.6, 0.3, "否定 / 证据不足", 12, True, AMBER_INK)
    text(slide, 6.92, 4.4, 5.6, 0.3, "“没有白毛，也不是晚疫病”", 13, True, AMBER_INK)
    text(slide, 6.92, 4.85, 5.6, 0.7, "→ 先追问发生部位、病斑形态和霉层\n→ 不生成用药条目", 11, False, AMBER_INK, line_spacing=1.3)
    text(slide, M, 6.62, 12.2, 0.3, "否定描述先剔除；候选接近或证据弱时澄清，而不是猜一个病名。", 10, False, MUTED)
    footer(slide, 8)

    # ---------------------------------------------------------------- 09 规则
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "08 · 校验", "逐条规则校验：把‘能说’变成‘能过’")
    rules = [
        ("01", "来源与身份", "编号、名称、剂型、证号一致"),
        ("02", "禁用清单", "命中即拦截"),
        ("03", "限用范围", "作物受限或范围不明"),
        ("04", "作物匹配", "依据记录必须同作物"),
        ("05", "防治对象", "依据记录必须同病害"),
        ("06", "剂量与单位", "同量纲换算，缺失不判过"),
        ("07", "间隔期与采收", "不得短于依据值"),
        ("08", "每季次数", "达到上限不再推荐"),
    ]
    for index, (number, heading, description) in enumerate(rules):
        col = index % 2
        row = index // 2
        x = M + col * 2.75
        y = 1.55 + row * 1.18
        rect(slide, x, y, 2.58, 1.0, WHITE, LINE, 0.08)
        text(slide, x + 0.12, y + 0.1, 0.35, 0.25, number, 10.5, True, AMBER)
        text(slide, x + 0.48, y + 0.1, 2.0, 0.25, heading, 11.2, True, PRIMARY)
        text(slide, x + 0.12, y + 0.46, 2.35, 0.45, description, 9.4, False, MUTED, line_spacing=1.2)

    audit = rounded_image(SHOTS / "agent-04-逐条校验.png", BUILD / "audit.png", 18)
    picture_card(slide, audit, 6.25, 1.55, 6.5)
    text(slide, 6.25, 5.35, 6.5, 0.3, "逐条展示：通过 / 待核验 / 拦截 + 原因", 10.5, True, PRIMARY)
    text(slide, 6.25, 5.75, 6.5, 0.8, "blocked 条目不推荐；needs_review 必须标注“待核验/需补充”；同一药剂的多个条目按索引分别裁决，不因同名串联。", 9.5, False, MUTED, line_spacing=1.3)
    footer(slide, 9)

    # ---------------------------------------------------------------- 10 窗口
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "09 · 窗口", "施药窗口：从‘能用什么’到‘什么时候能打’")
    flow = ["未来 72 小时逐时天气", "降雨 · 风速 · 温度 · 湿度", "清晨 06–09 / 傍晚 16–19", "按药剂倒推最晚施药日"]
    for index, label in enumerate(flow):
        y = 1.65 + index * 1.12
        rect(slide, M, y, 5.75, 0.72, WHITE if index < 3 else LIGHT, LINE, 0.1)
        text(slide, M + 0.2, y + 0.18, 5.35, 0.35, label, 11.5, index == 3, PRIMARY if index == 3 else INK)
        if index < 3:
            arrow(slide, 3.42, y + 0.75, 3.42, y + 1.08, "A9C9B7", 1.3)

    rect(slide, 6.75, 1.6, 6.0, 3.55, WHITE, LINE, 0.08)
    text(slide, 6.95, 1.75, 5.5, 0.3, "模拟窗口（以烯酰吗啉为例）", 12.5, True, PRIMARY)
    text(slide, 6.95, 2.2, 5.5, 0.55, "10/21  06:00–08:00", 21, True, GREEN)
    text(slide, 6.95, 2.78, 5.5, 0.3, "首选窗口", 10, False, MUTED)
    text(slide, 6.95, 3.18, 5.5, 0.5, "10/22  16:00–19:00", 16, True, GREEN)
    text(slide, 6.95, 3.66, 5.5, 0.3, "备选窗口 · 风速 2–3 级，建议降低喷头或加防风罩", 9.5, False, MUTED)
    text(slide, 6.95, 4.2, 5.5, 0.35, "最晚施药日 10/22", 13, True, AMBER_INK)
    text(slide, 6.95, 4.55, 5.5, 0.4, "窗口按每种药剂分别计算；已排除已开始的小时。", 9.5, False, MUTED)
    rect(slide, 11.55, 1.78, 1.0, 0.32, AMBER_BG, None, 0.2)
    text(slide, 11.55, 1.82, 1.0, 0.22, "模拟天气", 8.5, True, AMBER_INK, PP_ALIGN.CENTER)

    text(slide, M, 6.15, 12.2, 0.55, "窗口按每种药剂分别计算；过滤已开始的小时和超过最晚施药日的日期；不是当地预报。", 10.5, False, MUTED)
    footer(slide, 10)

    # ---------------------------------------------------------------- 11 安全边界
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "10 · 边界", "安全边界：把‘不知道’显式暴露出来")

    rect(slide, 3.4, 2.55, 6.55, 1.2, PRIMARY, None, 0.1)
    text(slide, 3.55, 2.7, 6.25, 0.45, "ready_for_field_use = false", 20, True, WHITE, PP_ALIGN.CENTER)
    text(slide, 3.55, 3.18, 6.25, 0.3, "竞赛演示，不签发正式田间处方", 10.5, False, "D7E8DD", PP_ALIGN.CENTER)

    cards = [
        (M, 1.55, "未提供次数 ≠ 0", "不把未知当零次"),
        (10.2, 1.55, "图文冲突先澄清", "不猜病名"),
        (M, 4.2, "纯色 / 模糊图片拒识", "推理前拦截"),
        (10.2, 4.2, "待核验只作演示", "不混入正式通过"),
    ]
    for x, y, heading, description in cards:
        rect(slide, x, y, 2.55, 1.0, WHITE, LINE, 0.08)
        text(slide, x + 0.15, y + 0.12, 2.25, 0.3, heading, 10.5, True, PRIMARY)
        text(slide, x + 0.15, y + 0.52, 2.25, 0.3, description, 9, False, MUTED)

    text(slide, M, 5.55, 6.0, 0.3, "知识库 40 条药剂记录", 10.5, True, PRIMARY)
    rect(slide, M, 5.95, 7.35, 0.42, LIGHT, None, 0.2)
    rect(slide, M, 5.95, 7.35 * 0.15, 0.42, GREEN, None, 0.2)
    text(slide, M + 7.5, 5.97, 4.6, 0.3, "6 条团队标准比对  ·  34 条待核验  ·  0 条现行登记核验", 9.5, False, MUTED)
    text(slide, M, 6.55, 12.2, 0.3, "标准比对不等于具体产品登记核验；待核验资料只能作为演示预览。", 9.5, False, MUTED)
    footer(slide, 11)

    # ---------------------------------------------------------------- 12 证据
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "11 · 证据", "受控数据上的可复现结果")

    metrics = [
        ("96.88%", "总体准确率"),
        ("6.53%", "测试集拒识率"),
        ("98.82%", "接受样本准确率"),
        ("36 / 36", "固定标签功能基准"),
    ]
    for index, (number, label) in enumerate(metrics):
        y = 1.7 + index * 1.18
        text(slide, M, y, 2.4, 0.55, number, 28, True, GREEN)
        text(slide, M + 2.45, y + 0.14, 2.2, 0.3, label, 11, False, MUTED)
    text(slide, M, 6.0, 6.0, 0.3, "60 项自动化测试全部通过", 11, True, GREEN)

    chart_data = CategoryChartData()
    chart_data.categories = ["接受样本", "拒识样本"]
    chart_data.add_series("测试集 1088 张", (1017, 71))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(7.1), Inches(1.65), Inches(5.4), Inches(4.2), chart_data
    )
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "测试集 1088 张"
    for paragraph in chart.chart_title.text_frame.paragraphs:
        for run in paragraph.runs:
            _set_font(run, 12, True, PRIMARY)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    chart.legend.font.name = FONT
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_percentage = True
    labels.show_value = True
    labels.font.size = Pt(9.5)
    labels.font.name = FONT
    labels.font.color.rgb = rgb(WHITE)
    series = plot.series[0]
    for point_index, color in enumerate([GREEN, AMBER]):
        point = series.points[point_index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
        point.format.line.fill.background()

    text(slide, 7.1, 6.0, 5.4, 0.3, "接受 1017 张 · 拒识 71 张；拒识率 6.53%", 9.5, False, MUTED)
    text(slide, M, 6.55, 12.2, 0.3, "来源：PlantVillage 公开数据集 + 本项目实跑；阈值由验证集确定；不是田间安全率。", 9, False, MUTED)
    footer(slide, 12)

    # ---------------------------------------------------------------- 13 创新点
    slide = prs.slides.add_slide(blank)
    bg(slide)
    title(slide, "12 · 创新", "让智能体‘能做事’，也让结果‘能核对’")
    rows = [
        ["维度", "通用问答 / 识别工具", "田诊卫士智能体"],
        ["依据", "只给答案，来源不稳定", "每条记录可回查，标注核验状态"],
        ["剂量与次数", "容易给出笼统建议", "逐条检查剂量、间隔期、每季次数"],
        ["不确定时", "可能猜一个答案", "先追问或拒识，不生成用药"],
        ["安全边界", "难以审计", "blocked 不推荐，ready_for_field_use=false"],
    ]
    add_table(slide, M, 1.55, 7.7, rows, [1.35, 2.75, 3.6], row_height=0.78, font_size=10.2, header_size=10.5)

    rect(slide, 8.55, 1.55, 4.2, 4.55, LIGHT, None, 0.08)
    text(slide, 8.78, 1.75, 3.75, 1.5, "“不是替代农技员，而是把依据和缺口摆到桌面上。”", 18, True, PRIMARY, line_spacing=1.35)
    bullets(
        slide,
        8.78,
        3.45,
        3.7,
        2.4,
        ["小农户：少一次盲目用药", "农技员：多一份可核对清单", "后续：更多作物、真实登记、真实气象"],
        10.5,
        gap=7,
    )
    text(slide, M, 6.55, 12.2, 0.3, "创新不只在模型，而在‘模型规划 + 代码否决 + 证据组织’的组合。", 10.5, False, MUTED)
    footer(slide, 13)

    # ---------------------------------------------------------------- 14 结尾
    slide = prs.slides.add_slide(blank)
    bg(slide, DARK)
    text(slide, 0.8, 0.75, 6.0, 0.3, "当前边界", 11.5, True, "A8D4BF")
    text(slide, 0.8, 1.15, 6.2, 1.2, "先把能做的做扎实，\n再把不能做的说清楚", 28, True, WHITE, line_spacing=1.28)
    bullets(
        slide,
        0.8,
        2.75,
        5.8,
        2.6,
        [
            "34 / 40 条药剂记录仍待核验",
            "天气为模拟数据，未接真实气象",
            "未开展田间试用与专家独立核验",
            "不签发正式田间处方",
        ],
        12.5,
        color="D7E8DD",
        bullet_color=AMBER,
        gap=9,
    )
    text(slide, 7.3, 0.75, 5.0, 0.3, "下一步", 11.5, True, "A8D4BF")
    bullets(
        slide,
        7.3,
        1.15,
        5.2,
        3.2,
        [
            "完成具体产品现行登记核验",
            "接入真实气象并核验当地规则",
            "小规模用户任务测试与田间试点",
        ],
        12.5,
        color="D7E8DD",
        bullet_color=AMBER,
        gap=9,
    )
    text(slide, 7.3, 4.65, 5.2, 0.9, "让创意被体验、\n被演示、能继续迭代。", 21, True, AMBER, line_spacing=1.3)
    text(slide, 0.8, 6.55, 7.0, 0.3, "田诊卫士 · 智能体版  ·  【队名】", 10.5, False, "A8D4BF")
    text(slide, 8.8, 6.55, 3.7, 0.3, "谢谢", 10.5, False, "A8D4BF", PP_ALIGN.RIGHT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"OK {OUT}  {OUT.stat().st_size / 1024 / 1024:.2f} MB")


if __name__ == "__main__":
    build()
